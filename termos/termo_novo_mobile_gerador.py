import streamlit as st
import os
import glob
import time
import locale
import platform
import getpass
import re
import shutil
import subprocess
from datetime import datetime
from docx import Document
import io

# Tenta importar sua classe de banco.
try:
    from Database import Database
except ImportError:
    st.error("Arquivo Database.py não encontrado. O sistema não funcionará corretamente.")
    class Database: pass

# Componente que dispara a cada tecla (busca automática ao completar a matrícula).
# Sem o pacote, o campo cai no modo padrão do Streamlit (busca no Enter).
try:
    from st_keyup import st_keyup
    KEYUP_DISPONIVEL = True
except ImportError:
    KEYUP_DISPONIVEL = False

# --- CONFIGURAÇÃO DA PÁGINA ---
st.set_page_config(
    page_title="Autorização Auto",
    page_icon="🖨️",
    layout="wide"
)

# Configuração de Localização
try:
    locale.setlocale(locale.LC_TIME, 'pt_BR.utf8')
except:
    try:
        locale.setlocale(locale.LC_TIME, 'Portuguese_Brazil.1252')
    except:
        locale.setlocale(locale.LC_TIME, 'C')

# --- ESTILOS CSS PERSONALIZADOS ---
st.markdown("""
    <style>
    .stButton>button {
        width: 100%;
        font-weight: bold;
    }
    .contexto-box {
        background-color: #f0f2f6;
        padding: 10px;
        border-radius: 5px;
        border-left: 5px solid #ff4b4b;
        margin-bottom: 10px;
    }
    </style>
""", unsafe_allow_html=True)

# --- CLASSES E UTILITÁRIOS ---

class SistemaAutorizacao:
    def __init__(self):
        self.dirs = {
            'templates': './templates',
            'output': './output',
            'logs_local': './logs',
            'logs_rede': r'\\192.168.10.3\Departamentos\Pública\autorizacaoAuto\logs'
        }
        
        self.tags_obrigatorias = [
            "{{RESPONSAVEL}}", "{{CPF RESPONSAVEL}}", "{{ENDERECO RESPONSAVEL}}",
            "{{NOME MENOR}}", "{{CPF MENOR}}", "{{IDADE MENOR}}",
            "{{DATA ASSINATURA}}", "{{MAT}}"
        ]

        self.tags_conhecidas = self.tags_obrigatorias + ["{{EMAIL RESPONSAVEL}}", "{{MAT}}", "{{IGNORAR}}"]
        
        # Inicializa pastas
        for d in [self.dirs['templates'], self.dirs['output'], self.dirs['logs_local']]:
            os.makedirs(d, exist_ok=True)

    def get_db_connection(self):
        return Database('ODBC Driver 18 for SQL Server', '192.168.10.12,56183', 'MultiClubes', 'lara', 'Mudar@123')

    def log(self, tipo, mensagem):
        user = getpass.getuser()
        machine = platform.node()
        timestamp = datetime.now().strftime('%d/%m/%Y %H:%M:%S')
        log_line = f"[{timestamp}] | {tipo:<10} | {user}@{machine} | {mensagem}\n"
        
        try:
            with open(os.path.join(self.dirs['logs_local'], f"log_{datetime.now().strftime('%Y-%m')}.txt"), 'a', encoding='utf-8') as f: f.write(log_line)
        except: pass

    def calcular_idade(self, data_nasc):
        if not data_nasc: return 0
        hoje = datetime.now()
        return hoje.year - data_nasc.year - ((hoje.month, hoje.day) < (data_nasc.month, data_nasc.day))

    def verificar_parentesco(self, nome_resp, nome_menor):
        ignorar = {'da', 'de', 'do', 'das', 'dos', 'e', 'junior', 'neto', 'filho', 'sobrinho', 'segundo', 'terceiro'}
        partes_resp = set(str(nome_resp).lower().split()) - ignorar
        partes_menor = set(str(nome_menor).lower().split()) - ignorar
        return bool(partes_resp & partes_menor)

    def validar_template(self, caminho_docx):
        try:
            doc = Document(caminho_docx)
            texto = " ".join([p.text for p in doc.paragraphs])
            for t in doc.tables:
                for r in t.rows:
                    for c in r.cells: texto += " " + c.text
            
            faltantes = [t for t in self.tags_obrigatorias if t not in texto]
            todas_tags = set(re.findall(r"\{\{.*?\}\}", texto))
            desconhecidas = [t for t in todas_tags if t not in self.tags_conhecidas]
            return faltantes, desconhecidas
        except Exception as e:
            return [str(e)], []

    def buscar_dados(self, mat):
        db = self.get_db_connection()
        try:
            t = db.select('Id, Street, Number, Burgh, City, State, PostalCode', 'dbo.Titles', f"dbo.Titles.Code = '{mat}'")
            if not t: return None
            
            end = ", ".join([str(x) for x in t[0][1:] if x])
            m = db.select("Name, Document, Rg, BirthDate, Email", "dbo.Members", f"dbo.Members.Title = '{t[0][0]}' AND dbo.Members.Status = 0")
            
            pessoas = []
            for p in m:
                idade = self.calcular_idade(p[3])
                pessoas.append({'nome': p[0], 'cpf': p[1], 'rg': p[2], 'idade': idade, 'email': p[4], 'endereco': end})
            return pessoas
        except Exception as e:
            st.error(f"Erro no banco de dados: {e}")
            return None

    def gerar_documento(self, template_path, resp, menor, mat=""):
        doc = Document(template_path)
        
        mapa = {
            "{{RESPONSAVEL}}": str(resp['nome']), 
            "{{CPF RESPONSAVEL}}": str(resp['cpf'] or "___.___.___-__"),
            "{{ENDERECO RESPONSAVEL}}": str(resp['endereco']), 
            "{{RG RESPONSAVEL}}": str(resp['rg'] or "______________"), 
            "{{RG MENOR}}": str(menor['rg'] or "______________"), 
            "{{EMAIL RESPONSAVEL}}": str(resp.get('email', '') or ""),
            "{{NOME MENOR}}": str(menor['nome']), 
            "{{CPF MENOR}}": str(menor['cpf'] or "___.___.___-__"),
            "{{IDADE MENOR}}": str(menor['idade']), 
            "{{DATA ASSINATURA}}": datetime.now().strftime('%d de %B de %Y'),
            "{{MAT}}": f"Matrícula: {mat}" if mat else ""
        }

        def processar_texto(obj):
            if hasattr(obj, 'text') and obj.text:
                for k,v in mapa.items(): 
                    if k in obj.text: obj.text = obj.text.replace(k,v)
                if "{{" in obj.text:
                    obj.text = re.sub(r"\{\{.*?\}\}", "__________________________", obj.text)

        for p in doc.paragraphs: processar_texto(p)
        for t in doc.tables:
            for r in t.rows:
                for c in r.cells:
                    for p in c.paragraphs: processar_texto(p)

        nome_arq = f"{str(resp['nome']).split()[0]}_{str(menor['nome']).split()[0]}.docx"
        path_final = os.path.abspath(os.path.join(self.dirs['output'], nome_arq))
        doc.save(path_final)
        self.log("GERACAO", f"Arquivo: {nome_arq} | Template: {os.path.basename(template_path)}")
        return path_final, nome_arq

    def localizar_soffice(self):
        """Localiza o executável do LibreOffice (soffice)."""
        candidatos = [
            r"C:\Program Files\LibreOffice\program\soffice.exe",
            r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
            "/usr/bin/soffice",
            "/opt/libreoffice/program/soffice",
        ]
        for c in candidatos:
            if os.path.exists(c):
                return c
        # Procura no PATH do sistema
        return shutil.which("soffice") or shutil.which("soffice.exe")

    def imprimir_servidor(self, path, impressora=None):
        """Imprime SEMPRE via LibreOffice (soffice), independente do app padrão do Windows."""
        soffice = self.localizar_soffice()
        if not soffice:
            msg = ("LibreOffice não encontrado. Instale o LibreOffice para imprimir "
                   "(esperado em 'C:\\Program Files\\LibreOffice\\program\\soffice.exe').")
            self.log("ERRO_IMP", msg)
            return False, msg

        # Perfil isolado para não conflitar com uma janela do LibreOffice já aberta.
        perfil = os.path.abspath(os.path.join(self.dirs['output'], '.lo_profile'))
        perfil_uri = 'file:///' + perfil.replace('\\', '/')

        cmd = [soffice, '--headless', '--norestore', '--nologo',
               f'-env:UserInstallation={perfil_uri}']
        if impressora:
            cmd += ['--pt', impressora, path]   # imprime na impressora informada
        else:
            cmd += ['-p', path]                 # imprime na impressora padrão do sistema

        try:
            resultado = subprocess.run(cmd, timeout=120, capture_output=True, text=True)
            self.log("IMPRESSAO",
                     f"LibreOffice -> {os.path.basename(path)} | "
                     f"impressora={impressora or 'PADRÃO'} | rc={resultado.returncode}")
            if resultado.returncode == 0:
                return True, f"Enviado para impressão via LibreOffice ({impressora or 'impressora padrão'})."
            erro = (resultado.stderr or resultado.stdout or "").strip()
            return False, f"LibreOffice retornou código {resultado.returncode}. {erro}"
        except subprocess.TimeoutExpired:
            self.log("ERRO_IMP", f"Timeout ao imprimir {os.path.basename(path)}")
            return False, "Tempo limite excedido ao enviar para o LibreOffice."
        except Exception as e:
            self.log("ERRO_IMP", f"{os.path.basename(path)} | {e}")
            return False, str(e)

    # --- LÓGICA DO GERADOR DE TEMPLATE ---
    def sugerir_tag(self, texto_anterior):
        """Tenta adivinhar a tag com base no texto que vem antes do underline"""
        txt = texto_anterior.lower()
        if "cpf" in txt and "menor" in txt: return "{{CPF MENOR}}"
        if "cpf" in txt: return "{{CPF RESPONSAVEL}}"
        if "menor" in txt or "criança" in txt: return "{{NOME MENOR}}"
        if "responsável" in txt or "autorizante" in txt: return "{{RESPONSAVEL}}"
        if "endereço" in txt or "residente" in txt: return "{{ENDERECO RESPONSAVEL}}"
        if "idade" in txt: return "{{IDADE MENOR}}"
        if "email" in txt: return "{{EMAIL RESPONSAVEL}}"
        if "data" in txt or "em" in txt: return "{{DATA ASSINATURA}}"
        return "{{IGNORAR}}"

    def criar_novo_template(self, doc_io, mapeamento, nome_arquivo):
        doc = Document(doc_io)
        
        # Regex para encontrar 3 ou mais underlines
        regex = re.compile(r'_{3,}')
        
        contador_subs = 0

        def substituir_no_objeto(obj):
            nonlocal contador_subs
            if hasattr(obj, 'text') and obj.text:
                # Se encontrar o padrão
                matches = regex.findall(obj.text)
                for match in matches:
                    if contador_subs < len(mapeamento):
                        tag_escolhida = mapeamento[contador_subs]
                        if tag_escolhida != "{{IGNORAR}}":
                            # Substitui apenas a primeira ocorrência deste match específico
                            obj.text = obj.text.replace(match, tag_escolhida, 1)
                        contador_subs += 1

        for p in doc.paragraphs: substituir_no_objeto(p)
        for t in doc.tables:
            for r in t.rows:
                for c in r.cells:
                    for p in c.paragraphs: substituir_no_objeto(p)
        
        caminho_salvar = os.path.join(self.dirs['templates'], nome_arquivo)
        if not caminho_salvar.endswith('.docx'): caminho_salvar += '.docx'
        doc.save(caminho_salvar)
        return caminho_salvar

# --- INICIALIZAÇÃO DO SISTEMA ---
app = SistemaAutorizacao()

# --- SIDEBAR: CONFIGURAÇÕES ---
st.sidebar.title("⚙️ Configurações")
st.sidebar.markdown("---")

template_files = glob.glob(os.path.join(app.dirs['templates'], "*.docx"))
template_names = [os.path.basename(x) for x in template_files]

selected_template_name = None
template_path = None

if template_names:
    selected_template_name = st.sidebar.selectbox("Modelo de Documento", template_names)
    template_path = os.path.join(app.dirs['templates'], selected_template_name)
    
    if st.sidebar.checkbox("Validar Modelo Atual"):
        faltantes, desconhecidas = app.validar_template(template_path)
        if faltantes: st.sidebar.error(f"Tags Faltando: {', '.join(faltantes)}")
        if desconhecidas: st.sidebar.warning(f"Tags Extras: {', '.join(desconhecidas)}")
        if not faltantes and not desconhecidas: st.sidebar.success("Modelo Válido!")
else:
    st.sidebar.warning("Nenhum template encontrado.")

st.sidebar.markdown("---")

# --- CONFIGURAÇÃO DE IMPRESSÃO (LibreOffice) ---
st.sidebar.subheader("🖨️ Impressão")
soffice_path = app.localizar_soffice()
if soffice_path:
    st.sidebar.success("LibreOffice detectado.")
else:
    st.sidebar.error("LibreOffice NÃO encontrado. A impressão não funcionará.")

impressora_nome = st.sidebar.text_input(
    "Impressora (opcional):",
    help="Deixe em branco para usar a impressora padrão do Windows. Informe o nome exato para forçar uma impressora específica."
).strip()
impressora_alvo = impressora_nome or None

st.sidebar.markdown("---")
st.sidebar.info(f"Servidor: {platform.node()}")

# --- ÁREA PRINCIPAL ---
st.title("🖨️ Emissor de Autorizações")

# Agora são 3 abas
tab1, tab2, tab3 = st.tabs(["👥 Sócio (Busca DB)", "📝 Não Sócio (Manual)", "🛠️ Gerador de Template"])

# --- TAB 1: SÓCIO ---
with tab1:
    if not template_path:
        st.error("Selecione um template na barra lateral para continuar.")
    else:
        TAM_MATRICULA = 5

        def executar_busca(matricula_digitada):
            """Consulta o banco e guarda o resultado no estado da sessão."""
            matricula_digitada = (matricula_digitada or "").strip()
            if not matricula_digitada:
                return
            # Evita reconsultar o banco para a mesma matrícula já carregada.
            if st.session_state.get('matricula_atual') == matricula_digitada and 'dados_socio' in st.session_state:
                return
            st.session_state.pop('dados_socio', None)
            st.session_state.pop('criancas_externas', None)
            dados = app.buscar_dados(matricula_digitada)
            if dados:
                st.session_state['dados_socio'] = dados
                st.session_state['matricula_atual'] = matricula_digitada
            else:
                st.session_state['erro_busca'] = f"Matrícula {matricula_digitada} não encontrada."

        col_search, col_btn = st.columns([3, 1])
        with col_search:
            if KEYUP_DISPONIVEL:
                matricula = st_keyup(
                    "Digite a Matrícula (busca automática ao completar 5 dígitos):",
                    max_chars=TAM_MATRICULA,
                    debounce=400,
                    key="mat_input"
                )
            else:
                matricula = st.text_input(
                    "Digite a Matrícula (Enter para buscar):",
                    max_chars=TAM_MATRICULA,
                    key="mat_input"
                )
        with col_btn:
            st.write("")
            st.write("")
            buscar_manual = st.button("🔍 Buscar", use_container_width=True)

        matricula = (matricula or "").strip()

        # Validação: matrícula deve conter apenas números.
        if matricula and not matricula.isdigit():
            st.warning("A matrícula deve conter apenas números.")
        else:
            st.session_state.pop('erro_busca', None)
            # Busca automática assim que os 5 dígitos são preenchidos,
            # ou manualmente pelo botão.
            if (len(matricula) == TAM_MATRICULA) or (buscar_manual and matricula):
                executar_busca(matricula)

        if st.session_state.get('erro_busca'):
            st.error(st.session_state['erro_busca'])

        if 'dados_socio' in st.session_state:
            dados = st.session_state['dados_socio']
            st.success(f"Matrícula {st.session_state['matricula_atual']} carregada.")
            
            adultos = [p for p in dados if p['idade'] >= 18]
            criancas = [p for p in dados if p['idade'] < 18]

            st.divider()
            
            col_resp, col_menor = st.columns(2)
            with col_resp:
                st.subheader("1. Responsável")
                if not adultos:
                    st.warning("Sem maiores de 18 anos.")
                    resp_selecionado = None
                else:
                    opcoes_adultos = {f"{p['nome']} ({p['idade']} anos)": p for p in adultos}
                    k_adulto = st.selectbox("Selecione o Responsável:", list(opcoes_adultos.keys()))
                    resp_selecionado = opcoes_adultos[k_adulto]
                    if resp_selecionado: st.info(f"CPF: {resp_selecionado['cpf']}")

            with col_menor:
                st.subheader("2. Menor")
                origem_menor = st.radio("Origem:", ["Mesma Matrícula", "Outra Matrícula"], horizontal=True)
                menor_selecionado = None
                
                if origem_menor == "Mesma Matrícula":
                    if not criancas:
                        st.warning("Sem menores nesta matrícula.")
                    else:
                        opcoes_criancas = {f"{p['nome']} ({p['idade']} anos)": p for p in criancas}
                        k_crianca = st.selectbox("Selecione o Menor:", list(opcoes_criancas.keys()))
                        menor_selecionado = opcoes_criancas[k_crianca]
                else: 
                    col_m_input, col_m_btn = st.columns([3,1])
                    with col_m_input: mat_menor_input = st.text_input("Matrícula do Menor:", max_chars=5)
                    with col_m_btn: 
                        st.write("")
                        st.write("")
                        btn_buscar_m = st.button("Buscar Menor")
                    
                    if btn_buscar_m and mat_menor_input:
                        dados_m2 = app.buscar_dados(mat_menor_input)
                        if dados_m2:
                            st.session_state['criancas_externas'] = [p for p in dados_m2 if p['idade'] < 18]
                        else: st.error("Matrícula não encontrada.")
                    
                    if 'criancas_externas' in st.session_state and st.session_state['criancas_externas']:
                        opcoes_c2 = {f"{p['nome']} ({p['idade']} anos)": p for p in st.session_state['criancas_externas']}
                        k_c2 = st.selectbox("Menor (Outra Mat.):", list(opcoes_c2.keys()))
                        menor_selecionado = opcoes_c2[k_c2]

            if resp_selecionado and menor_selecionado:
                st.divider()
                divergencia = not app.verificar_parentesco(resp_selecionado['nome'], menor_selecionado['nome'])
                check_ok = True
                if divergencia:
                    st.error("⚠️ ALERTA: Sobrenomes não conferem!")
                    st.write(f"**Resp:** {resp_selecionado['nome']} | **Menor:** {menor_selecionado['nome']}")
                    check_ok = st.checkbox("Estou ciente e desejo imprimir mesmo assim.")

                # Verificações adicionais: avisa sobre dados faltantes (não bloqueia).
                if not resp_selecionado.get('cpf'):
                    st.warning("Responsável sem CPF cadastrado — o campo ficará em branco no documento.")
                if not menor_selecionado.get('cpf'):
                    st.warning("Menor sem CPF cadastrado — o campo ficará em branco no documento.")
                if not resp_selecionado.get('endereco'):
                    st.warning("Endereço do responsável não encontrado no cadastro.")

                sem_impressora = app.localizar_soffice() is None
                if sem_impressora:
                    st.error("LibreOffice não localizado — instale-o para habilitar a impressão.")

                if st.button("🖨️ Gerar e Imprimir (Sócio)", type="primary",
                             disabled=(not check_ok) or sem_impressora):
                    path, nome = app.gerar_documento(template_path, resp_selecionado, menor_selecionado, st.session_state['matricula_atual'])
                    sucesso, msg = app.imprimir_servidor(path, impressora_alvo)
                    if sucesso:
                        st.balloons()
                        st.toast(msg, icon="🖨️")
                    else: st.error(f"Erro na impressão: {msg}")

# --- TAB 2: NÃO SÓCIO ---
with tab2:
    if not template_path:
        st.error("Selecione um template na barra lateral para continuar.")
    else:
        st.header("Preenchimento Manual")
        with st.form("form_nao_socio"):
            col1, col2 = st.columns(2)
            with col1:
                st.subheader("Responsável")
                ns_resp_nome = st.text_input("Nome Completo (Responsável)")
                ns_resp_cpf = st.text_input("CPF")
                ns_resp_end = st.text_input("Endereço Completo")
            with col2:
                st.subheader("Menor")
                ns_menor_nome = st.text_input("Nome Completo (Menor)")
                ns_menor_cpf = st.text_input("CPF Menor")
                ns_menor_idade = st.number_input("Idade", min_value=0, max_value=17, step=1)
            submitted = st.form_submit_button("Analisar e Gerar")
            if submitted:
                if not ns_resp_nome or not ns_menor_nome: st.error("Nomes são obrigatórios.")
                else:
                    divergencia = not app.verificar_parentesco(ns_resp_nome, ns_menor_nome)
                    if divergencia:
                        st.warning(f"⚠️ Sobrenomes diferentes: '{ns_resp_nome}' vs '{ns_menor_nome}'")
                        st.info("Para prosseguir, clique no botão de confirmação que aparecerá abaixo.")
                        st.session_state['confirmar_divergencia_ns'] = True
                    else:
                        st.session_state['confirmar_divergencia_ns'] = False
                        st.session_state['gerar_ns_trigger'] = True

        if st.session_state.get('confirmar_divergencia_ns'):
            if st.button("Confirmar divergência e Imprimir"):
                st.session_state['gerar_ns_trigger'] = True

        if st.session_state.get('gerar_ns_trigger'):
            r_obj = {'nome': ns_resp_nome, 'cpf': ns_resp_cpf, 'endereco': ns_resp_end, 'rg': '', 'email': ''}
            m_obj = {'nome': ns_menor_nome, 'cpf': ns_menor_cpf, 'idade': ns_menor_idade, 'rg': ''}
            path, nome = app.gerar_documento(template_path, r_obj, m_obj, "MANUAL")
            sucesso, msg = app.imprimir_servidor(path, impressora_alvo)
            if sucesso:
                st.balloons()
                st.success(f"Enviado para impressora: {nome}")
            else: st.error(msg)
            st.session_state['gerar_ns_trigger'] = False
            st.session_state['confirmar_divergencia_ns'] = False

# --- TAB 3: GERADOR DE TEMPLATE ---
with tab3:
    st.header("🛠️ Importar e Mapear Novo Modelo")
    st.markdown("Faça upload de um arquivo `.docx` com campos vazios (linhas `______`) e mapeie-os para as variáveis do sistema.")
    
    upload_file = st.file_uploader("Upload do Arquivo Word com linhas", type=['docx'])
    
    if upload_file:
        # Analisa o arquivo
        if 'analise_template' not in st.session_state or st.session_state.get('arquivo_atual') != upload_file.name:
            doc = Document(upload_file)
            regex = re.compile(r'_{3,}') # 3 ou mais underlines
            
            campos_encontrados = []
            
            def procurar_linhas(obj):
                if hasattr(obj, 'text') and obj.text:
                    for match in regex.finditer(obj.text):
                        # Pega 30 caracteres antes para contexto
                        start = max(0, match.start() - 30)
                        contexto = obj.text[start:match.start()].strip()
                        # Adiciona '...' se cortou texto
                        if start > 0: contexto = "..." + contexto
                        
                        sugestao = app.sugerir_tag(contexto)
                        
                        campos_encontrados.append({
                            'contexto': contexto,
                            'match': match.group(),
                            'sugestao': sugestao
                        })

            for p in doc.paragraphs: procurar_linhas(p)
            for t in doc.tables:
                for r in t.rows:
                    for c in r.cells:
                        for p in c.paragraphs: procurar_linhas(p)
            
            st.session_state['analise_template'] = campos_encontrados
            st.session_state['arquivo_atual'] = upload_file.name
            st.session_state['doc_bytes'] = upload_file.getvalue()

        # Exibe Interface de Mapeamento
        campos = st.session_state['analise_template']
        
        if not campos:
            st.warning("Nenhuma linha (_______) encontrada neste documento.")
        else:
            st.success(f"Encontrados {len(campos)} campos para preenchimento.")
            
            with st.form("form_template"):
                escolhas = []
                for i, campo in enumerate(campos):
                    st.markdown(f"<div class='contexto-box'><b>Campo {i+1}:</b> Texto anterior: <i>'{campo['contexto']}'</i></div>", unsafe_allow_html=True)
                    
                    # Selectbox com sugestão selecionada por padrão
                    try:
                        idx_sugestao = app.tags_conhecidas.index(campo['sugestao'])
                    except: idx_sugestao = 0
                    
                    escolha = st.selectbox(
                        f"O que colocar no Campo {i+1}?", 
                        app.tags_conhecidas,
                        index=idx_sugestao,
                        key=f"tag_{i}"
                    )
                    escolhas.append(escolha)
                    st.write("") # Espaço

                novo_nome = st.text_input("Nome para salvar o modelo (ex: termo_viagem):")
                salvar = st.form_submit_button("💾 Salvar Novo Template")
                
                if salvar:
                    if not novo_nome:
                        st.error("Digite um nome para o arquivo.")
                    else:
                        # Recupera o doc original da memória
                        doc_io = io.BytesIO(st.session_state['doc_bytes'])
                        
                        try:
                            caminho = app.criar_novo_template(doc_io, escolhas, novo_nome)
                            st.success(f"Template salvo com sucesso em: {caminho}")
                            st.info("Recarregue a página (F5) para ele aparecer na lista de templates.")
                            # Limpa estado
                            del st.session_state['analise_template']
                            del st.session_state['doc_bytes']
                        except Exception as e:
                            st.error(f"Erro ao salvar: {e}")