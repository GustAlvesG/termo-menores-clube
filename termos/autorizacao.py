"""Regras de negócio do Emissor de Autorizações (sem dependência de interface)."""
import getpass
import io
import os
import platform
import re
import shutil
import subprocess
import sys
import threading
import unicodedata
from datetime import datetime

from docx import Document
from dotenv import load_dotenv

# Pasta de dados: ao lado do .exe quando empacotado, ou ao lado deste arquivo.
if getattr(sys, "frozen", False):
    BASE_DIR = os.path.dirname(os.path.abspath(sys.executable))
else:
    BASE_DIR = os.path.dirname(os.path.abspath(__file__))

# Pasta raiz da Central de Impressão (um nível acima de termos/).
CENTRAL_DIR = os.path.dirname(BASE_DIR)

# Um .env local tem prioridade; o .env compartilhado da central completa o resto.
load_dotenv(os.path.join(BASE_DIR, ".env"))
load_dotenv(os.path.join(CENTRAL_DIR, ".env"))

# PDFs gerados pelo WhatsApp: só arquivos desta pasta podem ser impressos via API.
PASTA_PDFS_WHATSAPP = os.path.abspath(os.getenv(
    "WHATSAPP_PDF_DIR", os.path.join(CENTRAL_DIR, "whatsapp", "media", "para_impressao")))

MESES = ["janeiro", "fevereiro", "março", "abril", "maio", "junho", "julho",
         "agosto", "setembro", "outubro", "novembro", "dezembro"]

TAM_MATRICULA = 5
LINHA_EM_BRANCO = "__________________________"

TAGS_OBRIGATORIAS = [
    "{{RESPONSAVEL}}", "{{CPF RESPONSAVEL}}", "{{ENDERECO RESPONSAVEL}}",
    "{{NOME MENOR}}", "{{CPF MENOR}}", "{{IDADE MENOR}}",
    "{{DATA ASSINATURA}}", "{{MAT}}",
]
TAGS_CONHECIDAS = TAGS_OBRIGATORIAS + [
    "{{EMAIL RESPONSAVEL}}", "{{RG RESPONSAVEL}}", "{{RG MENOR}}", "{{IGNORAR}}",
]
# Grafias alternativas aceitas nos modelos (não aparecem no gerador de templates).
TAGS_ALIAS = {"{{E-MAIL RESPONSAVEL}}": "{{EMAIL RESPONSAVEL}}"}

RE_TAG = re.compile(r"\{\{.*?\}\}")
RE_LINHA = re.compile(r"_{3,}")


class ErroUsuario(Exception):
    """Erro com mensagem que pode ser mostrada diretamente ao usuário."""


def data_por_extenso(data=None):
    data = data or datetime.now()
    return f"{data.day} de {MESES[data.month - 1]} de {data.year}"


def nome_seguro(texto, padrao="arquivo"):
    """Remove acentos e caracteres inválidos para nome de arquivo."""
    texto = unicodedata.normalize("NFKD", str(texto)).encode("ascii", "ignore").decode()
    texto = re.sub(r"[^A-Za-z0-9_-]+", "_", texto).strip("_")
    return texto or padrao


def _paragrafos(doc):
    for p in doc.paragraphs:
        yield p
    for t in doc.tables:
        for r in t.rows:
            for c in r.cells:
                yield from c.paragraphs


class SistemaAutorizacao:
    def __init__(self):
        self.dirs = {
            "templates": os.path.join(BASE_DIR, "templates"),
            "output": os.path.join(BASE_DIR, "output"),
            "logs": os.getenv("LOGS_DIR") or os.path.join(BASE_DIR, "logs"),
        }
        for d in self.dirs.values():
            os.makedirs(d, exist_ok=True)
        # O LibreOffice usa um perfil compartilhado: duas impressões simultâneas
        # nele falham, então as impressões são enfileiradas.
        self._lock_impressao = threading.Lock()

    # --- infraestrutura ---
    def get_db_connection(self):
        from Database import Database
        faltando = [v for v in ("DB_SERVER", "DB_NAME", "DB_USER", "DB_PASSWORD") if not os.getenv(v)]
        if faltando:
            raise ErroUsuario(f"Configuração do banco ausente no .env: {', '.join(faltando)}")
        return Database(
            os.getenv("DB_DRIVER", "ODBC Driver 18 for SQL Server"),
            os.getenv("DB_SERVER"), os.getenv("DB_NAME"),
            os.getenv("DB_USER"), os.getenv("DB_PASSWORD"),
        )

    def log(self, tipo, mensagem, origem=None):
        origem = origem or f"{getpass.getuser()}@{platform.node()}"
        timestamp = datetime.now().strftime("%d/%m/%Y %H:%M:%S")
        linha = f"[{timestamp}] | {tipo:<10} | {origem} | {mensagem}\n"
        try:
            arq = os.path.join(self.dirs["logs"], f"log_{datetime.now():%Y-%m}.txt")
            with open(arq, "a", encoding="utf-8") as f:
                f.write(linha)
        except OSError:
            pass

    # --- templates ---
    def listar_templates(self):
        return sorted(f for f in os.listdir(self.dirs["templates"])
                      if f.lower().endswith(".docx") and not f.startswith("~$"))

    def caminho_template(self, nome):
        # Só aceita nomes que existem na pasta (impede acesso a outros caminhos).
        if nome not in self.listar_templates():
            raise ErroUsuario("Modelo de documento não encontrado.")
        return os.path.join(self.dirs["templates"], nome)

    def validar_template(self, nome):
        doc = Document(self.caminho_template(nome))
        texto = " ".join(p.text for p in _paragrafos(doc))
        faltantes = [t for t in TAGS_OBRIGATORIAS if t not in texto]
        desconhecidas = sorted(t for t in set(RE_TAG.findall(texto))
                               if t not in TAGS_CONHECIDAS and t not in TAGS_ALIAS)
        return faltantes, desconhecidas

    # --- dados ---
    @staticmethod
    def calcular_idade(data_nasc):
        if not data_nasc:
            return None
        hoje = datetime.now()
        return hoje.year - data_nasc.year - ((hoje.month, hoje.day) < (data_nasc.month, data_nasc.day))

    @staticmethod
    def verificar_parentesco(nome_resp, nome_menor):
        ignorar = {"da", "de", "do", "das", "dos", "e", "junior", "neto", "filho",
                   "sobrinho", "segundo", "terceiro"}
        partes_resp = set(str(nome_resp).lower().split()) - ignorar
        partes_menor = set(str(nome_menor).lower().split()) - ignorar
        return bool(partes_resp & partes_menor)

    @staticmethod
    def validar_matricula(mat):
        mat = (mat or "").strip()
        if not mat.isdigit() or len(mat) > TAM_MATRICULA:
            raise ErroUsuario(f"A matrícula deve ter até {TAM_MATRICULA} dígitos numéricos.")
        return mat

    def buscar_dados(self, mat):
        """Retorna a lista de membros ativos da matrícula, ou None se não existir."""
        mat = self.validar_matricula(mat)
        db = self.get_db_connection()
        t = db.select("Id, Street, Number, Burgh, City, State, PostalCode",
                      "dbo.Titles", "dbo.Titles.Code = ?", (mat,))
        if not t:
            return None
        endereco = ", ".join(str(x) for x in t[0][1:] if x)
        membros = db.select("Name, Document, Rg, BirthDate, Email", "dbo.Members",
                            "dbo.Members.Title = ? AND dbo.Members.Status = 0 ORDER BY Name",
                            (t[0][0],))
        return [{
            "nome": m[0], "cpf": m[1] or "", "rg": m[2] or "",
            "idade": self.calcular_idade(m[3]), "email": m[4] or "",
            "endereco": endereco,
        } for m in membros]

    # --- documento ---
    def gerar_documento(self, nome_template, resp, menor, mat=""):
        doc = Document(self.caminho_template(nome_template))
        mapa = {
            "{{RESPONSAVEL}}": str(resp.get("nome", "")),
            "{{CPF RESPONSAVEL}}": str(resp.get("cpf") or "___.___.___-__"),
            "{{ENDERECO RESPONSAVEL}}": str(resp.get("endereco") or LINHA_EM_BRANCO),
            "{{RG RESPONSAVEL}}": str(resp.get("rg") or "______________"),
            "{{EMAIL RESPONSAVEL}}": str(resp.get("email") or LINHA_EM_BRANCO),
            "{{NOME MENOR}}": str(menor.get("nome", "")),
            "{{CPF MENOR}}": str(menor.get("cpf") or "___.___.___-__"),
            "{{RG MENOR}}": str(menor.get("rg") or "______________"),
            "{{IDADE MENOR}}": "" if menor.get("idade") is None else str(menor["idade"]),
            "{{DATA ASSINATURA}}": data_por_extenso(),
            "{{MAT}}": f"Matrícula: {mat}" if mat else "",
        }
        for alias, tag in TAGS_ALIAS.items():
            mapa[alias] = mapa[tag]

        def substituir(texto):
            for k, v in mapa.items():
                texto = texto.replace(k, v)
            return RE_TAG.sub(LINHA_EM_BRANCO, texto)

        for p in _paragrafos(doc):
            if "{{" not in p.text:
                continue
            # Primeiro tenta trocar dentro de cada trecho (preserva negrito/fonte).
            for run in p.runs:
                if "{{" in run.text:
                    run.text = substituir(run.text)
            # Se alguma tag ficou dividida entre trechos, troca no parágrafo inteiro.
            if "{{" in p.text:
                p.text = substituir(p.text)

        nome_arq = (f"{nome_seguro(str(resp.get('nome', '')).split()[0])}_"
                    f"{nome_seguro(str(menor.get('nome', '')).split()[0])}_"
                    f"{datetime.now():%Y%m%d_%H%M%S_%f}.docx")
        caminho = os.path.join(self.dirs["output"], nome_arq)
        doc.save(caminho)
        self.log("GERACAO", f"Arquivo: {nome_arq} | Template: {nome_template}")
        return caminho, nome_arq

    # --- impressão ---
    @staticmethod
    def localizar_soffice():
        candidatos = [
            r"C:\Program Files\LibreOffice\program\soffice.exe",
            r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
            "/usr/bin/soffice",
            "/opt/libreoffice/program/soffice",
        ]
        for c in candidatos:
            if os.path.exists(c):
                return c
        return shutil.which("soffice") or shutil.which("soffice.exe")

    def imprimir(self, path, impressora=None):
        """Imprime via LibreOffice. Retorna (sucesso, mensagem)."""
        soffice = self.localizar_soffice()
        if not soffice:
            msg = "LibreOffice não encontrado no servidor. Instale-o para imprimir."
            self.log("ERRO_IMP", msg)
            return False, msg

        perfil = os.path.abspath(os.path.join(self.dirs["output"], ".lo_profile"))
        perfil_uri = "file:///" + perfil.replace("\\", "/")
        cmd = [soffice, "--headless", "--norestore", "--nologo",
               f"-env:UserInstallation={perfil_uri}"]
        cmd += ["--pt", impressora, path] if impressora else ["-p", path]

        alvo = impressora or "impressora padrão"
        try:
            with self._lock_impressao:
                r = subprocess.run(cmd, timeout=120, capture_output=True, text=True)
            self.log("IMPRESSAO", f"{os.path.basename(path)} | impressora={alvo} | rc={r.returncode}")
            if r.returncode == 0:
                return True, f"Enviado para {alvo}."
            erro = (r.stderr or r.stdout or "").strip()
            return False, f"LibreOffice retornou código {r.returncode}. {erro}"
        except subprocess.TimeoutExpired:
            self.log("ERRO_IMP", f"Timeout ao imprimir {os.path.basename(path)}")
            return False, "Tempo limite excedido ao enviar para a impressora."
        except Exception as e:
            self.log("ERRO_IMP", f"{os.path.basename(path)} | {e}")
            return False, str(e)

    @staticmethod
    def localizar_sumatra():
        """SumatraPDF que vem junto do pacote pdf-to-printer usado pelo WhatsApp."""
        candidatos = [
            os.getenv("SUMATRA_PDF", ""),
            os.path.join(CENTRAL_DIR, "whatsapp", "node_modules", "pdf-to-printer",
                         "dist", "SumatraPDF-3.4.6-32.exe"),
        ]
        return next((c for c in candidatos if c and os.path.exists(c)), None)

    def imprimir_pdf(self, path, impressora=None, origem=None):
        """Imprime um PDF do WhatsApp na mesma fila dos termos. Retorna (sucesso, mensagem)."""
        path = os.path.abspath(path or "")
        if os.path.commonpath([path, PASTA_PDFS_WHATSAPP]) != PASTA_PDFS_WHATSAPP:
            raise ErroUsuario("Arquivo fora da pasta de PDFs do WhatsApp.")
        if not path.lower().endswith(".pdf") or not os.path.isfile(path):
            raise ErroUsuario("PDF não encontrado.")
        sumatra = self.localizar_sumatra()
        if not sumatra:
            msg = "SumatraPDF não encontrado (rode 'npm install' em whatsapp/)."
            self.log("ERRO_IMP", msg, origem)
            return False, msg

        cmd = [sumatra] + (["-print-to", impressora] if impressora else ["-print-to-default"])
        cmd += ["-silent", path]
        alvo = impressora or "impressora padrão"
        try:
            with self._lock_impressao:
                r = subprocess.run(cmd, timeout=120, capture_output=True, text=True)
            self.log("IMPRESSAO", f"{os.path.basename(path)} | impressora={alvo} | rc={r.returncode}", origem)
            if r.returncode == 0:
                return True, f"Enviado para {alvo}."
            return False, f"SumatraPDF retornou código {r.returncode}. {(r.stderr or r.stdout or '').strip()}"
        except subprocess.TimeoutExpired:
            self.log("ERRO_IMP", f"Timeout ao imprimir {os.path.basename(path)}", origem)
            return False, "Tempo limite excedido ao enviar para a impressora."
        except Exception as e:
            self.log("ERRO_IMP", f"{os.path.basename(path)} | {e}", origem)
            return False, str(e)

    # --- gerador de template ---
    @staticmethod
    def sugerir_tag(texto_anterior):
        txt = texto_anterior.lower()
        if "cpf" in txt and "menor" in txt: return "{{CPF MENOR}}"
        if "cpf" in txt: return "{{CPF RESPONSAVEL}}"
        if "rg" in txt and "menor" in txt: return "{{RG MENOR}}"
        if re.search(r"\brg\b", txt): return "{{RG RESPONSAVEL}}"
        if "menor" in txt or "criança" in txt: return "{{NOME MENOR}}"
        if "responsável" in txt or "autorizante" in txt or re.search(r"\beu\b", txt): return "{{RESPONSAVEL}}"
        if "endereço" in txt or "residente" in txt: return "{{ENDERECO RESPONSAVEL}}"
        if "idade" in txt: return "{{IDADE MENOR}}"
        if "email" in txt or "e-mail" in txt: return "{{EMAIL RESPONSAVEL}}"
        if "data" in txt or re.search(r"\bem\s*$", txt): return "{{DATA ASSINATURA}}"
        return "{{IGNORAR}}"

    def analisar_linhas(self, conteudo_docx):
        """Encontra as linhas em branco (___) do documento e sugere uma tag para cada."""
        doc = Document(io.BytesIO(conteudo_docx))
        campos = []
        for p in _paragrafos(doc):
            fim_anterior = 0
            for m in RE_LINHA.finditer(p.text):
                # Contexto = texto entre a linha anterior e esta (até 40 caracteres).
                inicio = max(fim_anterior, m.start() - 40)
                contexto = p.text[inicio:m.start()].strip()
                if inicio > fim_anterior:
                    contexto = "..." + contexto
                fim_anterior = m.end()
                campos.append({"contexto": contexto, "sugestao": self.sugerir_tag(contexto)})
        return campos

    def criar_novo_template(self, conteudo_docx, mapeamento, nome):
        invalidas = [t for t in mapeamento if t not in TAGS_CONHECIDAS]
        if invalidas:
            raise ErroUsuario(f"Tags inválidas: {', '.join(invalidas)}")
        doc = Document(io.BytesIO(conteudo_docx))
        indice = 0

        def trocar(m):
            nonlocal indice
            tag = mapeamento[indice] if indice < len(mapeamento) else "{{IGNORAR}}"
            indice += 1
            return m.group() if tag == "{{IGNORAR}}" else tag

        for p in _paragrafos(doc):
            if RE_LINHA.search(p.text):
                p.text = RE_LINHA.sub(trocar, p.text)

        nome_arq = nome_seguro(re.sub(r"\.docx$", "", nome.strip(), flags=re.I), "") + ".docx"
        if nome_arq == ".docx":
            raise ErroUsuario("Nome do modelo inválido.")
        doc.save(os.path.join(self.dirs["templates"], nome_arq))
        self.log("TEMPLATE", f"Novo modelo salvo: {nome_arq}")
        return nome_arq
